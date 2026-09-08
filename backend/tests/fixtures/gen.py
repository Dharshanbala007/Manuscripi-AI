"""In-memory DOCX/ZIP builders for tests. Keeps fixtures small and explicit."""

from __future__ import annotations

import io
import zipfile

from docx import Document


def minimal_docx_bytes(text: str = "Hello world") -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def docx_bytes_from(build) -> bytes:
    """`build(doc)` receives a python-docx Document to populate; returns the bytes."""
    doc = Document()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def zip_bytes(entries: dict[str, bytes] | None = None, *, include_required: bool = True) -> bytes:
    """Build an arbitrary ZIP. With `include_required`, adds the two parts a DOCX must have."""
    entries = dict(entries or {})
    if include_required:
        entries.setdefault("[Content_Types].xml", b"<?xml version='1.0'?><Types/>")
        entries.setdefault("word/document.xml", b"<?xml version='1.0'?><document/>")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, content in entries.items():
            zf.writestr(name, content)
    return buf.getvalue()


__all__ = ["docx_bytes_from", "minimal_docx_bytes", "zip_bytes"]
