import pytest
from docx import Document
from docx.shared import Inches
from docx.table import Table

from app.formatting.body import restyle_body
from app.formatting.figures import check_figures
from app.formatting.frontmatter import rebuild_frontmatter
from app.formatting.page import available_column_width_emu
from app.formatting.styles import ensure_styles
from app.formatting.tables import style_tables
from app.parsing.ooxml import iter_block_items
from app.profiles.loader import load_profile
from tests.fixtures.gen import academic_docx, docx_with, manuscript_from, write_docx

PROFILE = load_profile("ieee")


def test_restyle_body_applies_heading_and_body_styles(tmp_path):
    path = write_docx(tmp_path, academic_docx())
    ms = manuscript_from(path)
    doc = Document(str(path))
    ensure_styles(doc, PROFILE)

    changes = restyle_body(doc, ms, PROFILE)

    paragraph_styles = {it.style.name for it in iter_block_items(doc) if not isinstance(it, Table)}
    assert "MS Heading 1" in paragraph_styles
    assert "MS Body" in paragraph_styles
    assert any("heading" in c.lower() for c in changes)
    assert any("paragraph" in c.lower() for c in changes)


def test_restyle_body_keeps_inline_bold_emphasis(tmp_path):
    builder = Document()
    title = builder.add_paragraph("A Short Working Title")
    title.add_run("").bold = True
    builder.add_heading("1 Introduction", level=1)
    para = builder.add_paragraph("This sentence has a ")
    para.add_run("bold").bold = True
    para.add_run(" word and then keeps going for a while so it is clearly a body paragraph.")
    path = tmp_path / "in.docx"
    builder.save(str(path))

    ms = manuscript_from(path)
    doc = Document(str(path))
    ensure_styles(doc, PROFILE)

    restyle_body(doc, ms, PROFILE)

    paras = [it for it in iter_block_items(doc) if not isinstance(it, Table)]
    body_para = paras[-1]
    assert body_para.style.name == "MS Body"
    assert any(r.bold for r in body_para.runs)  # emphasis preserved


def test_rebuild_frontmatter_uses_edited_metadata(tmp_path):
    path = write_docx(tmp_path, academic_docx())
    ms = manuscript_from(path)
    ms.metadata.title.set_by_user("An Edited Title For The Paper")
    doc = Document(str(path))
    ensure_styles(doc, PROFILE)

    changes = rebuild_frontmatter(doc, ms, PROFILE)

    head = "\n".join(p.text for p in doc.paragraphs[:8])
    assert "An Edited Title For The Paper" in head
    assert any("title" in c.lower() for c in changes)


def test_style_tables_flags_wide_table(tmp_path):
    path = write_docx(tmp_path, docx_with(paragraphs=["intro"], table=(2, 12)))
    ms = manuscript_from(path)
    doc = Document(str(path))
    issues = style_tables(doc, ms, PROFILE)
    assert any("column" in i.message.lower() for i in issues)


def test_style_tables_shrinks_a_wide_table_to_fit_the_column(tmp_path):
    builder = Document()
    table = builder.add_table(rows=1, cols=3)
    for col, width in zip(table.columns, [Inches(3), Inches(2), Inches(1)], strict=True):
        col.width = width  # 6in total, well over IEEE's ~3.5in column
    path = tmp_path / "in.docx"
    builder.save(str(path))

    ms = manuscript_from(path)
    doc = Document(str(path))
    style_tables(doc, ms, PROFILE)

    avail = available_column_width_emu(PROFILE)
    resized = list(doc.tables[0].columns)
    total = sum(int(c.width) for c in resized)
    assert total <= avail
    assert total > avail * 0.99  # scaled to fill the column, not shrunk further than needed
    ratio = int(resized[0].width) / int(resized[1].width)
    assert ratio == pytest.approx(3 / 2, rel=1e-3)  # proportions preserved


def test_style_tables_leaves_a_narrow_table_untouched(tmp_path):
    builder = Document()
    table = builder.add_table(rows=1, cols=2)
    widths = [Inches(0.75), Inches(0.75)]
    for col, width in zip(table.columns, widths, strict=True):
        col.width = width  # 1.5in total, already well under the column
    path = tmp_path / "in.docx"
    builder.save(str(path))

    ms = manuscript_from(path)
    doc = Document(str(path))
    style_tables(doc, ms, PROFILE)

    assert [int(c.width) for c in doc.tables[0].columns] == [int(w) for w in widths]


def test_check_figures_flags_missing_caption(tmp_path):
    path = write_docx(tmp_path, docx_with(paragraphs=["some text"], image=True))
    ms = manuscript_from(path)
    doc = Document(str(path))
    issues = check_figures(doc, ms, PROFILE)
    assert any("caption" in i.message.lower() for i in issues)
