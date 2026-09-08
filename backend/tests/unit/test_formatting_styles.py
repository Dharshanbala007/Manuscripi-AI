from docx import Document
from docx.oxml.ns import qn

from app.formatting.page import apply_page_layout
from app.formatting.styles import ensure_styles
from app.profiles.loader import load_profile

_EXPECTED_STYLES = [
    "MS Title",
    "MS Author",
    "MS Affiliation",
    "MS Abstract",
    "MS Keywords",
    "MS Heading 1",
    "MS Heading 2",
    "MS Heading 3",
    "MS Body",
    "MS Caption",
    "MS Reference",
]


def test_ensure_styles_creates_named_styles():
    doc = Document()
    profile = load_profile("ieee")
    ensure_styles(doc, profile)

    names = {s.name for s in doc.styles}
    for expected in _EXPECTED_STYLES:
        assert expected in names

    body = doc.styles["MS Body"]
    assert body.font.name == "Times New Roman"
    assert round(body.font.size.pt, 1) == 10.0

    h1 = doc.styles["MS Heading 1"]
    assert abs(h1.font.size.pt - profile.headings.levels[1].size_pt) < 0.01


def test_ensure_styles_is_idempotent():
    doc = Document()
    profile = load_profile("ieee")
    ensure_styles(doc, profile)
    count = len(doc.styles)
    ensure_styles(doc, profile)
    assert len(doc.styles) == count


def test_apply_page_layout_sets_geometry_and_columns():
    doc = Document()
    profile = load_profile("ieee")
    apply_page_layout(doc, profile)

    section = doc.sections[0]
    assert abs(section.left_margin.inches - profile.page.margin_left_in) < 0.01
    assert abs(section.page_width.inches - 8.5) < 0.01

    cols = section._sectPr.find(qn("w:cols"))
    assert cols is not None
    assert cols.get(qn("w:num")) == "2"


def test_apply_page_layout_is_idempotent_on_columns():
    doc = Document()
    profile = load_profile("ieee")
    apply_page_layout(doc, profile)
    apply_page_layout(doc, profile)
    cols = doc.sections[0]._sectPr.findall(qn("w:cols"))
    assert len(cols) == 1
