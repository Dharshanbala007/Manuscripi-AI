"""Apply the profile table style, fit table width to the column, and flag tables whose
column count still makes them cramped even after fitting."""

from __future__ import annotations

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Twips

from app.domain.elements import ElementType
from app.domain.issues import Issue, IssueCategory, Severity
from app.formatting.page import available_column_width_emu
from app.profiles.base import PublisherProfile


def style_tables(doc, manuscript, profile: PublisherProfile) -> list[Issue]:
    issues: list[Issue] = []
    table_blocks = [b for b in manuscript.body if b.kind == ElementType.TABLE]
    avail_width = available_column_width_emu(profile)

    for i, table in enumerate(doc.tables):
        try:
            table.style = doc.styles[profile.tables.style_name]
        except KeyError:
            pass

        if profile.tables.header_bold and len(table.rows) > 0:
            for cell in table.rows[0].cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.bold = True

        _fit_table_width(table, avail_width)

        ncols = _column_count(table)
        if ncols > profile.validation.max_table_cols_before_flag:
            issues.append(
                Issue(
                    id=f"tbl-{i}-wide",
                    severity=Severity.WARNING,
                    category=IssueCategory.LAYOUT,
                    message=(
                        f"Table {i + 1} has {ncols} columns and may overflow the "
                        f"{profile.columns.count}-column layout."
                    ),
                    location=table_blocks[i].id if i < len(table_blocks) else None,
                    suggested_action="Split the table or place it across the full page width.",
                )
            )

    return issues


def _column_count(table) -> int:
    try:
        return len(table.columns)
    except Exception:  # noqa: BLE001 - irregular grid
        return len(table.rows[0].cells) if table.rows else 0


def _fit_table_width(table, avail_emu: int) -> None:
    """Scale an explicitly-sized table down to fit the profile's column width, keeping
    each column's proportion. Cell text is never touched -- a narrower column just wraps
    onto more lines -- and a table that already fits, or has no explicit width, is left
    alone (this only prevents overflow, it never enlarges a table)."""
    if table._tbl.tblPr is None:  # every real-world table has one; nothing safe to resize
        return
    current = _explicit_width_emu(table)
    if not current or current <= avail_emu:
        return
    scale = avail_emu / current
    try:
        cols = list(table.columns)
    except Exception:  # noqa: BLE001 - irregular grid; nothing safe to resize
        return
    if not all(c.width for c in cols):
        return
    for col in cols:
        col.width = Emu(round(int(col.width) * scale))
    _set_tblw_dxa(table, avail_emu)
    table.autofit = False


def _explicit_width_emu(table) -> int | None:
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW")) if tblPr is not None else None
    if tblW is not None and tblW.get(qn("w:type")) == "dxa":
        return int(Twips(int(tblW.get(qn("w:w")))))
    try:
        widths = [c.width for c in table.columns]
    except Exception:  # noqa: BLE001 - irregular grid
        return None
    return sum(int(w) for w in widths) if widths and all(w for w in widths) else None


def _set_tblw_dxa(table, width_emu: int) -> None:
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(int(Emu(width_emu).twips)))
