"""Generate the sample manuscripts used by tests and the UI.

Usage:  python scripts/gen_samples.py

Writes sample_documents/sample_{basic,complex,messy}.docx. Content is original,
deliberately synthetic academic filler - no copyrighted text.
"""

from __future__ import annotations

import base64
import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt

_ROOT = Path(__file__).resolve().parents[1]
_OUT = _ROOT / "sample_documents"

_PNG_1x1 = base64.b64decode(
    b"iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)

_LOREM = (
    "The proposed approach organises the manuscript into an ordered stream of "
    "elements and applies rule-based heuristics to classify each one. Results "
    "are reported across several synthetic datasets and configurations."
)


def _title(doc: Document, text: str, size: int = 20) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _equation(doc: Document, lead: str) -> None:
    p = doc.add_paragraph(lead)
    p._p.append(parse_xml(f"<m:oMath {nsdecls('m')}><m:r><m:t>E=mc^2</m:t></m:r></m:oMath>"))


def _table(doc: Document, n: int, rows: int = 3, cols: int = 3) -> None:
    t = doc.add_table(rows=rows, cols=cols)
    for c in range(cols):
        t.cell(0, c).text = f"Column {c + 1}"
    for r in range(1, rows):
        for c in range(cols):
            t.cell(r, c).text = f"{r}.{c}"
    doc.add_paragraph(f"Table {n}. Summary of experiment {n} parameters.")


def _figure(doc: Document, n: int, *, caption: bool = True) -> None:
    doc.add_picture(io.BytesIO(_PNG_1x1))
    if caption:
        doc.add_paragraph(f"Figure {n}. Overview of stage {n} of the pipeline.")


def build_basic() -> Document:
    doc = Document()
    _title(doc, "A Rule-Based Approach to Academic Manuscript Formatting")
    doc.add_paragraph("Ada Lovelace, Alan Turing and Grace Hopper")
    doc.add_paragraph("Department of Computing, University of Somewhere")
    doc.add_paragraph("ada@uni.example")

    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(
        "This paper presents a rule based system for detecting and normalising the "
        "structure of academic manuscripts. " + _LOREM
    )
    doc.add_paragraph("Index Terms— formatting, parsing, classification, reproducibility")

    for i, (name, body) in enumerate(
        [
            ("Introduction", f"Academic manuscripts follow shared conventions [1]. {_LOREM} [2]"),
            ("Related Work", f"Earlier systems handled layout only [3]. {_LOREM} [4]"),
            ("Methodology", f"We parse the document into an ordered stream. {_LOREM}"),
            ("Results", f"The system preserved content in every trial [5]. {_LOREM}"),
            ("Conclusion", f"Rule-based processing is sufficient for reliable output [6]."),
        ],
        start=1,
    ):
        doc.add_heading(f"{i}. {name}", level=1)
        doc.add_paragraph(body)
        if name == "Methodology":
            _table(doc, 1)
        if name == "Results":
            _figure(doc, 1)

    doc.add_heading("References", level=1)
    for i in range(1, 7):
        doc.add_paragraph(f"[{i}] Author {i}, A relevant work number {i}, 20{10 + i}.")
    return doc


def build_complex() -> Document:
    doc = Document()
    _title(doc, "Structured Detection and Reformatting of Multi-Section Research Manuscripts", 22)
    doc.add_paragraph("Ada Lovelace¹, Alan Turing², Grace Hopper¹ and Katherine Johnson³")
    doc.add_paragraph("¹ Department of Computing, University of Somewhere")
    doc.add_paragraph("² Institute for Advanced Study, Elsewhere")
    doc.add_paragraph("³ Numerical Analysis Laboratory, Farfield College")
    doc.add_paragraph("ada@uni.example, alan@ias.example, grace@uni.example")

    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(
        "We describe a rule-based pipeline that classifies manuscript elements, detects "
        "canonical sections, and applies a configurable publisher profile. " + _LOREM
    )
    doc.add_paragraph("Index Terms— document engineering, classification, layout, IEEE, Springer")

    sections = [
        ("Introduction", 3),
        ("Related Work", 3),
        ("Materials and Methods", 4),
        ("Experimental Setup", 2),
        ("Results", 4),
        ("Discussion", 3),
        ("Conclusion", 2),
    ]
    cite = 1
    fig_n = 1
    tbl_n = 1
    for idx, (name, paras) in enumerate(sections, start=1):
        doc.add_heading(f"{idx}. {name}", level=1)
        for j in range(paras):
            doc.add_paragraph(f"{_LOREM} [{cite}] and further discussion [{cite + 1}].")
            cite += 2
        if name in ("Materials and Methods", "Experimental Setup", "Results"):
            doc.add_heading(f"{idx}.1 Detail", level=2)
            doc.add_paragraph(f"{_LOREM} [{cite}]")
            cite += 1
            _equation(doc, "The governing relation is given by")
        if name == "Materials and Methods":
            _table(doc, tbl_n)
            tbl_n += 1
            _figure(doc, fig_n)
            fig_n += 1
        if name == "Results":
            _table(doc, tbl_n)
            tbl_n += 1
            _table(doc, tbl_n)
            tbl_n += 1
            _figure(doc, fig_n)
            fig_n += 1
            _figure(doc, fig_n)
            fig_n += 1
            _figure(doc, fig_n)
            fig_n += 1

    total_refs = max(25, cite)
    doc.add_heading("References", level=1)
    for i in range(1, total_refs + 1):
        doc.add_paragraph(f"[{i}] Author {i} and Coauthor {i}, Study number {i} on the topic, 20{i % 25:02d}.")
    return doc


def build_messy() -> Document:
    doc = Document()
    doc.add_paragraph(
        "In this manuscript we consider a broad question about formatting and we begin "
        "immediately without a clear title line above this sentence."
    )
    doc.add_paragraph("A. Researcher and B. Collaborator")
    doc.add_paragraph("some.email@nowhere.example")

    doc.add_paragraph(
        "This work looks at manuscript structure. " + _LOREM + " It has no Abstract heading."
    )
    doc.add_paragraph("Keywords: messy input, ambiguous title, missing caption")

    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(f"{_LOREM} [1] and also [2].")

    doc.add_paragraph("BACKGROUND AND MOTIVATION")  # all-caps ad-hoc heading
    doc.add_paragraph(f"{_LOREM} [3]")

    doc.add_heading("3 Methods", level=1)  # inconsistent numbering, missing dot
    doc.add_paragraph(f"{_LOREM} [4] [5]")
    _figure(doc, 1, caption=False)  # figure with no caption
    _figure(doc, 2, caption=True)

    doc.add_heading("Findings", level=1)
    doc.add_paragraph(f"{_LOREM} [6] [7]")

    doc.add_heading("References", level=1)
    for i in range(1, 8):
        doc.add_paragraph(f"[{i}] Author {i}, Work {i}, 20{10 + i}.")
    doc.add_paragraph("[8] Uncited Author, A work never referenced in the body, 2099.")
    return doc


def main() -> None:
    _OUT.mkdir(parents=True, exist_ok=True)
    for name, builder in (
        ("sample_basic", build_basic),
        ("sample_complex", build_complex),
        ("sample_messy", build_messy),
    ):
        path = _OUT / f"{name}.docx"
        builder().save(str(path))
        print(f"wrote {path.relative_to(_ROOT)}")


if __name__ == "__main__":
    main()
